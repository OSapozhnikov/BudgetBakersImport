from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.errors import AppError, ErrorMessage
from app.services.counterparty import extract_counterparty

# Some bank exports use Latin "C" instead of Cyrillic "С" in "Статус" / "Cтатус".
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "status": ("статус", "cтатус"),
    "date": ("дата операції", "дата"),
    "note": ("опис операції", "опис"),
    "card": ("рахунок/картка", "рахунок", "картка"),
    "category": ("категорія",),
    "amount": ("сума",),
    "currency": ("валюта",),
}

COMPLETED_STATUS = "виконано"

CURRENCY_ALIASES: dict[str, str] = {
    "₴": "UAH",
    "грн": "UAH",
    "грн.": "UAH",
    "uah": "UAH",
    "usd": "USD",
    "$": "USD",
    "eur": "EUR",
    "€": "EUR",
}

# Excel serial date epoch (Windows 1900 date system).
_EXCEL_EPOCH = datetime(1899, 12, 30)


@dataclass
class ParsedRow:
    date: date
    note: str
    card: str
    bank_category: str
    amount: Decimal
    currency: str
    status: str
    source_row: int
    counter_party: str = ""


@dataclass
class ParseResult:
    rows: list[ParsedRow] = field(default_factory=list)
    skipped: int = 0
    bank_categories: list[str] = field(default_factory=list)
    warnings: list[ErrorMessage] = field(default_factory=list)


def _norm_header(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\xa0", " ").strip().lower()


def _find_columns(headers: list[Any]) -> dict[str, int]:
    normalized = [_norm_header(h) for h in headers]
    found: dict[str, int] = {}
    for key, aliases in COLUMN_ALIASES.items():
        for idx, header in enumerate(normalized):
            if not header:
                continue
            if header in aliases or any(header.startswith(a) for a in aliases):
                found[key] = idx
                break
    required = ("date", "note", "category", "amount", "currency")
    missing = [k for k in required if k not in found]
    if missing:
        raise AppError(
            "parse.missing_columns",
            missing=", ".join(missing),
            headers=headers,
        )
    return found


def _parse_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if _is_blank(value):
        raise AppError("parse.empty_date")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return (_EXCEL_EPOCH + timedelta(days=float(value))).date()
        except (OverflowError, ValueError) as exc:
            raise AppError("parse.bad_date", value=repr(value)) from exc
    text = str(value).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    raise AppError("parse.bad_date", value=repr(value))


def _parse_amount(value: Any) -> Decimal:
    if _is_blank(value):
        raise AppError("parse.empty_amount")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return Decimal(str(value))
    text = (
        str(value)
        .strip()
        .replace("\xa0", "")
        .replace(" ", "")
        .replace(",", ".")
    )
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise AppError("parse.bad_amount", value=repr(value)) from exc


def normalize_currency(value: Any) -> str:
    if _is_blank(value):
        return "UAH"
    raw = str(value).strip()
    key = raw.lower()
    return CURRENCY_ALIASES.get(key, CURRENCY_ALIASES.get(raw, raw.upper()))


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and not str(value).strip():
        return True
    return False


def _cell_str(value: Any) -> str:
    if _is_blank(value):
        return ""
    return str(value).strip()


def _is_completed(status: Any) -> bool:
    if _is_blank(status):
        return True
    text = str(status).strip().splitlines()[0].strip().lower()
    return text == COMPLETED_STATUS


def parse_excel(content: bytes) -> ParseResult:
    """Parse bank Excel export (BudgetBakers-compatible column layout) into rows."""
    bio = BytesIO(content)
    wb = load_workbook(bio, read_only=True, data_only=True)
    try:
        ws = wb.active
        if ws is None:
            return ParseResult(warnings=[ErrorMessage("parse.empty_file")])

        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if header_row is None:
            return ParseResult(warnings=[ErrorMessage("parse.empty_file")])

        columns = list(header_row)
        if all(_is_blank(c) for c in columns):
            return ParseResult(warnings=[ErrorMessage("parse.empty_file")])

        colmap = _find_columns(columns)
        result = ParseResult()
        categories: set[str] = set()
        has_data = False

        for idx, values in enumerate(rows_iter):
            has_data = True
            source_row = idx + 2  # header is row 1
            cells = list(values)
            status_val = cells[colmap["status"]] if "status" in colmap and colmap["status"] < len(cells) else None
            if not _is_completed(status_val):
                result.skipped += 1
                continue
            try:
                op_date = _parse_date(_at(cells, colmap["date"]))
                amount = _parse_amount(_at(cells, colmap["amount"]))
                currency = normalize_currency(_at(cells, colmap["currency"]))
                note_s = _cell_str(_at(cells, colmap["note"]))
                cat_s = _cell_str(_at(cells, colmap["category"]))
                card_s = _cell_str(_at(cells, colmap["card"])) if "card" in colmap else ""
                status_s = (
                    str(status_val).strip().splitlines()[0].strip()
                    if not _is_blank(status_val)
                    else ""
                )
            except AppError as exc:
                result.warnings.append(
                    ErrorMessage(
                        "parse.row_skipped",
                        {"source_row": source_row, "reason": exc.code, **exc.params},
                    )
                )
                result.skipped += 1
                continue

            if cat_s:
                categories.add(cat_s)
            result.rows.append(
                ParsedRow(
                    date=op_date,
                    note=note_s,
                    card=card_s,
                    bank_category=cat_s,
                    amount=amount,
                    currency=currency,
                    status=status_s,
                    source_row=source_row,
                    counter_party=extract_counterparty(note_s),
                )
            )

        result.bank_categories = sorted(categories)
        if not has_data and not result.rows:
            result.warnings.append(ErrorMessage("parse.empty_file"))
        elif not result.rows:
            result.warnings.append(ErrorMessage("parse.no_completed_rows"))
        return result
    finally:
        wb.close()


def _at(cells: list[Any], index: int) -> Any:
    if index < 0 or index >= len(cells):
        return None
    return cells[index]
